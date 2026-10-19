from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
exec(open('items.py').read())
f=0.44
items=[(*a[:5],int(round(a[5]*f/100)*100),*a[6:]) for a in items]

AR="Arial"
H=Font(name=AR,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3864")
B=Font(name=AR,bold=True); N=Font(name=AR); BL=Font(name=AR,color="0000FF"); GR=Font(name=AR,color="008000")
Y=PatternFill("solid",fgColor="FFF2CC"); TOT=PatternFill("solid",fgColor="D9E1F2")
thin=Side(style="thin",color="BFBFBF"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
USD='#,##0.00;(#,##0.00);"-"'; USD0='#,##0;(#,##0);"-"'; PCT='0.0%;(0.0%);"-"'; KG='#,##0;(#,##0);"-"'

wb=Workbook()
# ---------- Параметры
p=wb.active; p.title="Параметры"
p["A1"]="Параметры расчёта импорта дверной фурнитуры (Китай → Бишкек)"; p["A1"].font=Font(name=AR,bold=True,size=14)
p["A2"]="Жёлтые ячейки с синим шрифтом — вводные, их можно менять; все листы пересчитаются."; p["A2"].font=Font(name=AR,italic=True,color="595959")
params=[
("Курс USD → KGS (сом)",87.5,'#,##0.00',"Ориентир; подставьте курс НБКР на дату оплаты"),
("Стоимость фуры Китай → Бишкек, $",9000,USD0,"Со слов заказчика"),
("Грузоподъёмность фуры, кг",20000,KG,"Типично 20–22 т для тента 13,6 м; фурнитура тяжёлая — упрёмся в вес, а не в объём"),
("Полезный объём фуры, м³",86,'#,##0',"Тент 13,6×2,45×2,6 м"),
("Доля фрахта до границы ЕАЭС (входит в там. стоимость)",0.8,PCT,"Таможенная стоимость = цена + доставка до места ввоза (Иркештам/Торугарт). Основной путь — по Китаю, поэтому ~80%"),
("Страхование груза, % от FOB",0.003,'0.00%',"Ориентир 0,2–0,5%"),
("Таможенный сбор за оформление, % от там. стоимости",0.004,'0.00%',"0,40% — разъяснение ГТС КР (economist.kg, 03.04.2025)"),
("НДС при импорте, %",0.12,PCT,"Ставка НДС КР 12%"),
("Включать НДС в себестоимость? (1 = да, 0 = нет)",1,'0',"Если вы плательщик НДС и берёте его в зачёт — поставьте 0"),
("Таможенный брокер, декларация, $",500,USD0,"Ориентир за 1 ДТ с 20 позициями"),
("СВХ, разгрузка, доставка до склада, $",400,USD0,"Ориентир"),
("Комиссия банка / платёжного агента, % от FOB",0.01,PCT,"Перевод в Китай; 0,5–2%"),
("Агент / байер в Китае, % от FOB",0.03,PCT,"Поиск фабрик, консолидация, контроль; 3–5%. Если работаете напрямую — 0"),
("Образцы, инспекция качества, $",500,USD0,"Ориентир"),
("Резерв на непредвиденные, %",0.03,PCT,"Простой на границе, пересорт, курсовая разница"),
("Оптовая наценка, %",0.35,PCT,"Для оценки выручки; поменяйте под ваш рынок"),
]
p["A4"]="Параметр"; p["B4"]="Значение"; p["C4"]="Комментарий / источник"
for c in "ABC": p[c+"4"].font=H; p[c+"4"].fill=HF
R={}
for i,(n,v,fmt,c) in enumerate(params):
    r=5+i; p[f"A{r}"]=n; p[f"B{r}"]=v; p[f"C{r}"]=c
    p[f"A{r}"].font=N; p[f"C{r}"].font=Font(name=AR,color="595959")
    p[f"B{r}"].font=BL; p[f"B{r}"].fill=Y; p[f"B{r}"].number_format=fmt
    for col in "ABC": p[f"{col}{r}"].border=BOX
    R[i]=f"Параметры!$B${r}"
RATE,TRUCK,CAPKG,CAPM3,SHARE,INS,FEE,VAT,VATIN,BROKER,SVH,BANK,AGENT,INSP,CONT,MARK=[R[i] for i in range(16)]
p.column_dimensions["A"].width=54; p.column_dimensions["B"].width=14; p.column_dimensions["C"].width=95

# ---------- Заказ
z=wb.create_sheet("Заказ")
z["A1"]="Заказ на одну фуру и себестоимость по позициям, USD"; z["A1"].font=Font(name=AR,bold=True,size=14)
z["A2"]="Синие — вводные (кол-во, цена FOB, вес, объём, ставка пошлины). Цены — ориентир по made-in-china / Global Sources, уточните у фабрик. Логистика распределена пропорционально весу."
z["A2"].font=Font(name=AR,italic=True,color="595959")
hdr=["№","Категория","Наименование / спецификация","Код ТН ВЭД","Пошлина","Ед.","Кол-во","Цена FOB, $","Сумма FOB, $",
"Вес ед. брутто, кг","Вес итого, кг","Объём ед., м³","Объём итого, м³","Логистика, $","Страховка, $","Там. стоимость, $",
"Пошлина, $","Там. сбор, $","НДС импорт, $","Прочие расходы, $","Резерв, $","Себестоимость итого, $",
"Себест. ед., $","Себест. ед., сом","Опт. цена ед., сом","Себест. / FOB"]
HR=4
for j,h in enumerate(hdr,1):
    c=z.cell(HR,j,h); c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center"); c.border=BOX
z.row_dimensions[HR].height=45
r0=HR+1; r1=r0+len(items)-1; T=r1+1
fixed=f"({BROKER}+{SVH}+{INSP})"
for i,(cat,name,hs,duty,unit,qty,fob,kg,m3) in enumerate(items):
    r=r0+i
    vals={1:i+1,2:cat,3:name,4:hs,5:duty,6:unit,7:qty,8:fob,10:kg,12:m3}
    for k,v in vals.items(): z.cell(r,k,v)
    z[f"I{r}"]=f"=G{r}*H{r}"
    z[f"K{r}"]=f"=G{r}*J{r}"
    z[f"M{r}"]=f"=G{r}*L{r}"
    z[f"N{r}"]=f"=IF($K${T}=0,0,K{r}/$K${T}*{TRUCK})"
    z[f"O{r}"]=f"=I{r}*{INS}"
    z[f"P{r}"]=f"=I{r}+N{r}*{SHARE}+O{r}"
    z[f"Q{r}"]=f"=P{r}*E{r}"
    z[f"R{r}"]=f"=P{r}*{FEE}"
    z[f"S{r}"]=f"=(P{r}+Q{r})*{VAT}"
    z[f"T{r}"]=f"=IF($I${T}=0,0,I{r}/$I${T}*{fixed})+I{r}*({BANK}+{AGENT})"
    z[f"U{r}"]=f"=(I{r}+N{r}+O{r}+Q{r}+R{r}+S{r}*{VATIN}+T{r})*{CONT}"
    z[f"V{r}"]=f"=I{r}+N{r}+O{r}+Q{r}+R{r}+S{r}*{VATIN}+T{r}+U{r}"
    z[f"W{r}"]=f"=IF(G{r}=0,0,V{r}/G{r})"
    z[f"X{r}"]=f"=W{r}*{RATE}"
    z[f"Y{r}"]=f"=X{r}*(1+{MARK})"
    z[f"Z{r}"]=f"=IF(H{r}=0,0,W{r}/H{r})"
    for col in range(1,27):
        c=z.cell(r,col); c.border=BOX
        c.font = BL if col in (5,7,8,10,12) else N
        if col in (5,7,8,10,12): c.fill=Y
    for col,fmt in {5:'0%',7:KG,8:USD,9:USD0,10:'0.00',11:KG,12:'0.0000',13:'0.0',14:USD0,15:USD0,16:USD0,17:USD0,18:USD0,19:USD0,20:USD0,21:USD0,22:USD0,23:USD,24:'#,##0',25:'#,##0',26:'0.00"x"'}.items():
        z.cell(r,col).number_format=fmt
z[f"C{T}"]="ИТОГО"
for col in "GIKMNOPQRSTUV":
    z[f"{col}{T}"]=f"=SUM({col}{r0}:{col}{r1})"
z[f"Z{T}"]=f"=IF(I{T}=0,0,V{T}/I{T})"
for col in range(1,27):
    c=z.cell(T,col); c.font=B; c.fill=TOT; c.border=BOX
    c.number_format={7:KG,11:KG,13:'0.0',26:'0.00"x"'}.get(col,USD0)
# capacity check
z[f"C{T+2}"]="Загрузка фуры по весу"; z[f"G{T+2}"]=f"=K{T}/{CAPKG}"
z[f"C{T+3}"]="Загрузка фуры по объёму"; z[f"G{T+3}"]=f"=M{T}/{CAPM3}"
z[f"C{T+4}"]="Проверка"; z[f"G{T+4}"]=f'=IF(OR(G{T+2}>1,G{T+3}>1),"ПЕРЕГРУЗ — уменьшите кол-во","Помещается в одну фуру")'
for rr in (T+2,T+3): z[f"G{rr}"].number_format=PCT
for rr in (T+2,T+3,T+4): z[f"C{rr}"].font=B; z[f"G{rr}"].font=B
widths=[4,12,52,15,9,7,9,10,11,9,10,9,9,11,10,12,10,9,10,11,9,13,10,10,10,9]
for j,w in enumerate(widths,1): z.column_dimensions[L(j)].width=w
z.freeze_panes="D5"
z["E4"].comment=Comment("Ставки ЕТТ ЕАЭС: 8302 41 — 10%, 8302 10 (петли) — 0%, 8302 60 (доводчики) — 10%, 8301 40 — 12%, 8301 60 (цилиндры как части замков) — 15%. Источники: alta.ru, tws.by, ifcg.ru","Claude")

# ---------- Итог
s=wb.create_sheet("Итог",0)
s["A1"]="Себестоимость партии дверной фурнитуры — одна фура из Китая"; s["A1"].font=Font(name=AR,bold=True,size=14)
s["A3"]="Статья"; s["B3"]="USD"; s["C3"]="сом"; s["D3"]="% к FOB"
for c in "ABCD": s[c+"3"].font=H; s[c+"3"].fill=HF
lines=[("Товар (FOB Китай)",f"=Заказ!I{T}"),("Логистика (фура)",f"=Заказ!N{T}"),("Страхование",f"=Заказ!O{T}"),
("Ввозная пошлина",f"=Заказ!Q{T}"),("Таможенный сбор 0,4%",f"=Заказ!R{T}"),
("НДС при импорте (в себестоимости, если включён)",f"=Заказ!S{T}*{VATIN}"),
("Брокер, СВХ, банк, агент, инспекция",f"=Заказ!T{T}"),("Резерв на непредвиденные",f"=Заказ!U{T}")]
for i,(n,fm) in enumerate(lines):
    r=4+i; s[f"A{r}"]=n; s[f"B{r}"]=fm; s[f"C{r}"]=f"=B{r}*{RATE}"; s[f"D{r}"]=f"=IF($B$4=0,0,B{r}/$B$4)"
    s[f"B{r}"].font=GR; s[f"A{r}"].font=N
rt=4+len(lines)
s[f"A{rt}"]="СЕБЕСТОИМОСТЬ ПАРТИИ"; s[f"B{rt}"]=f"=SUM(B4:B{rt-1})"; s[f"C{rt}"]=f"=B{rt}*{RATE}"; s[f"D{rt}"]=f"=IF($B$4=0,0,B{rt}/$B$4)"
for r in range(4,rt+1):
    s[f"B{r}"].number_format=USD0; s[f"C{r}"].number_format='#,##0'; s[f"D{r}"].number_format=PCT
    for c in "ABCD": s[f"{c}{r}"].border=BOX
for c in "ABCD": s[f"{c}{rt}"].font=B; s[f"{c}{rt}"].fill=TOT
s[f"A{rt+1}"]="Справочно: НДС к уплате на границе (деньги нужны в любом случае)"; s[f"B{rt+1}"]=f"=Заказ!S{T}"; s[f"B{rt+1}"].number_format=USD0; s[f"C{rt+1}"]=f"=B{rt+1}*{RATE}"; s[f"C{rt+1}"].number_format='#,##0'
s[f"A{rt+2}"]="Нужно денег всего (товар + все расходы + НДС)"; s[f"B{rt+2}"]=f"=B{rt}+Заказ!S{T}*(1-{VATIN})"; s[f"B{rt+2}"].number_format=USD0; s[f"C{rt+2}"]=f"=B{rt+2}*{RATE}"; s[f"C{rt+2}"].number_format='#,##0'
for c in "AB": s[f"{c}{rt+2}"].font=B

k=rt+4
s[f"A{k}"]="Показатели"; s[f"A{k}"].font=H; s[f"A{k}"].fill=HF; s[f"B{k}"].fill=HF
kp=[("Вес партии, кг",f"=Заказ!K{T}",KG),("Загрузка фуры по весу",f"=Заказ!K{T}/{CAPKG}",PCT),
("Объём партии, м³",f"=Заказ!M{T}",'0.0'),("Загрузка фуры по объёму",f"=Заказ!M{T}/{CAPM3}",PCT),
("Единиц товара в заказе, шт/компл",f"=Заказ!G{T}",KG),
("Логистика на 1 кг, $",f"=Заказ!N{T}/Заказ!K{T}",USD),
("Таможенные платежи (пошлина+сбор+НДС), $",f"=Заказ!Q{T}+Заказ!R{T}+Заказ!S{T}",USD0),
("Себестоимость к FOB (коэффициент)",f"=B{rt}/B4",'0.00"x"'),
("Выручка при оптовой наценке, $",f"=B{rt}*(1+{MARK})",USD0),
("Валовая прибыль с партии, $",f"=B{rt}*{MARK}",USD0)]
for i,(n,fm,fmt) in enumerate(kp):
    r=k+1+i; s[f"A{r}"]=n; s[f"B{r}"]=fm; s[f"B{r}"].number_format=fmt; s[f"A{r}"].font=N
    for c in "AB": s[f"{c}{r}"].border=BOX
# by category
c0=k+len(kp)+3
s[f"A{c0}"]="По категориям"; s[f"B{c0}"]="FOB, $"; s[f"C{c0}"]="Себестоимость, $"; s[f"D{c0}"]="Вес, кг"
for c in "ABCD": s[f"{c}{c0}"].font=H; s[f"{c}{c0}"].fill=HF
cats=[]
for a in items:
    if a[0] not in cats: cats.append(a[0])
for i,cat in enumerate(cats):
    r=c0+1+i; s[f"A{r}"]=cat
    s[f"B{r}"]=f'=SUMIF(Заказ!$B${r0}:$B${r1},A{r},Заказ!$I${r0}:$I${r1})'
    s[f"C{r}"]=f'=SUMIF(Заказ!$B${r0}:$B${r1},A{r},Заказ!$V${r0}:$V${r1})'
    s[f"D{r}"]=f'=SUMIF(Заказ!$B${r0}:$B${r1},A{r},Заказ!$K${r0}:$K${r1})'
    s[f"B{r}"].number_format=USD0; s[f"C{r}"].number_format=USD0; s[f"D{r}"].number_format=KG
    for c in "ABCD": s[f"{c}{r}"].border=BOX; s[f"{c}{r}"].font=N
for col,w in zip("ABCD",[62,16,18,12]): s.column_dimensions[col].width=w

# ---------- Источники
src=wb.create_sheet("Источники и допущения")
rows=[("Что","Источник / допущение"),
("Пошлина 8302 41 (ручки, защёлки, фиксаторы, шпингалеты для дверей) — 10%","alta.ru/tnved/code/8302419000 ; tws.by/tws/tnved/code/10389"),
("Пошлина 8301 40 (замки врезные) — 12%","ifcg.ru/kb/tnved/8301401100 ; ifcg.ru/kb/tnved/8301401900"),
("Пошлина 8302 10 (петли) — 0%; 8302 60 000 9 (доводчики) — 10%; 8301 60 000 9 (цилиндры как части замков) — 15%","tws.by/tws/tnved/code/10385 ; /10397 ; /10383"),
("Таможенный сбор 0,40% от там. стоимости","Разъяснение ГТС КР, economist.kg, 03.04.2025"),
("НДС при импорте в КР — 12%","Ставка НДС КР; режим условного начисления изменён пост. №565 от 20.08.2026 (economist.kg)"),
("Ручки на розетке, цинк: $1,8–4 за компл., премиум $5–10","made-in-china.com (Handle Trends, Door Lock Handle), Global Sources"),
("Корпус врезного замка: от $1,56 (нерж., MOQ 600), 0,65 кг","globalsources.com, Lock body 8545C-3"),
("Петли, цилиндры, защёлки, шпингалеты, доводчики","Типовые FOB-цены фабрик Вэньчжоу/Гуандуна; уточнить по RFQ"),
("Фура $9 000","Со слов заказчика"),
("Вес и объём ед. — с упаковкой (брутто)","Оценка по типовым экспортным коробкам; уточнить по упаковочным листам фабрик"),
("Сертификация ЕАЭС","В расчёт не включена — уточнить у брокера, нужна ли декларация/сертификат по позициям"),
]
for i,(a,b) in enumerate(rows):
    src.cell(i+1,1,a); src.cell(i+1,2,b)
    for c in (1,2):
        src.cell(i+1,c).font = H if i==0 else N
        if i==0: src.cell(i+1,c).fill=HF
src.column_dimensions["A"].width=80; src.column_dimensions["B"].width=90
wb.save("Импорт_фурнитуры_Китай_расчёт.xlsx")
